"""Importer #1 (P0): FSANZ Australian Food Composition Database, Release 2.

Operating model: an operator downloads the AFCD workbook from
foodstandards.gov.au (no crawling required) and runs:

    python -m macromatch_ingest.cli afcd path/to/afcd_release2.xlsx

AFCD publishes nutrient values per 100 g edible portion with stable
"Public Food Key" identifiers — ideal for idempotent upserts. Column headers
have shifted slightly between releases, so mapping is done by fuzzy header
aliases rather than hard-coded column letters; unmapped required columns fail
loudly at startup instead of importing garbage.

Supports .xlsx (openpyxl) and .csv exports of the same table.
"""
from __future__ import annotations

import csv
import logging
from pathlib import Path
from typing import Iterator, Optional

from ..models import RawRecord, normalise_text
from ..units import to_float
from .base import BaseImporter

log = logging.getLogger("macromatch.afcd")

# header aliases -> canonical nutrient keys (values per 100 g)
COLUMN_ALIASES: dict[str, tuple[str, ...]] = {
    "external_id": ("public food key", "food key"),
    "name": ("food name", "food description"),
    "category": ("classification name", "classification", "food group"),
    "energy_kj": ("energy with dietary fibre equated kj", "energy with dietary fibre kj",
                  "energy kj"),
    "protein_g": ("protein g",),
    "fat_g": ("fat total g", "total fat g"),
    "carbs_g": ("available carbohydrate with sugar alcohols g",
                "available carbohydrates with sugar alcohols g",
                "carbohydrate g", "available carbohydrate g"),
    "sugars_g": ("total sugars g", "sugars g"),
    "sat_fat_g": ("saturated fatty acids g", "total saturated fatty acids g",
                  "saturated fat g"),
    "sodium_mg": ("sodium na mg", "sodium mg"),
    "fibre_g": ("total dietary fibre g", "dietary fibre g"),
}
REQUIRED_COLUMNS = ("external_id", "name", "energy_kj", "protein_g", "fat_g", "carbs_g")


def _map_headers(headers: list[str]) -> dict[str, int]:
    """Map canonical keys -> column index using normalised alias matching."""
    norm = [normalise_text(h) for h in headers]
    mapping: dict[str, int] = {}
    for key, aliases in COLUMN_ALIASES.items():
        for alias in aliases:
            target = normalise_text(alias)
            for i, h in enumerate(norm):
                if h == target or h.startswith(target):
                    mapping[key] = i
                    break
            if key in mapping:
                break
    missing = [k for k in REQUIRED_COLUMNS if k not in mapping]
    if missing:
        raise ValueError(
            f"AFCD file is missing required columns {missing}; found headers: {headers[:12]}...")
    return mapping


def _iter_rows(path: Path) -> Iterator[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        for row in ws.iter_rows(values_only=True):
            yield ["" if c is None else c for c in row]
    elif path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            yield from csv.reader(fh)
    else:
        raise ValueError(f"unsupported AFCD file type: {path.suffix}")


class AFCDImporter(BaseImporter):
    source_id = "afcd_r2"
    source_version = "AFCD Release 2.0"
    SOURCE_URL = "https://www.foodstandards.gov.au/science-data/monitoringnutrients/afcd"

    def iter_raw(self, input_ref: str, resume_after: Optional[str]) -> Iterator[tuple[str, RawRecord]]:
        path = Path(input_ref)
        if not path.exists():
            raise FileNotFoundError(
                f"{path} not found — download the AFCD workbook from {self.SOURCE_URL} first")
        rows = _iter_rows(path)
        # find the header row (AFCD sheets sometimes carry a title row above headers)
        mapping: Optional[dict[str, int]] = None
        headers: list[str] = []
        for probe in range(5):
            try:
                headers = [str(c) for c in next(rows)]
            except StopIteration:
                break
            try:
                mapping = _map_headers(headers)
                break
            except ValueError:
                continue
        if mapping is None:
            raise ValueError("could not locate AFCD header row in the first 5 rows")
        log.info("AFCD columns mapped: %s", {k: headers[v] for k, v in mapping.items()})

        def cell(row: list, key: str):
            i = mapping.get(key)
            return row[i] if i is not None and i < len(row) else None

        for n, row in enumerate(rows):
            cursor = f"{n:08d}"
            if resume_after and cursor <= resume_after:
                continue
            external_id = str(cell(row, "external_id") or "").strip()
            name = str(cell(row, "name") or "").strip()
            if not external_id or not name:
                continue  # blank/summary rows
            raw = {headers[i]: row[i] for i in range(min(len(headers), len(row)))}
            raw["_fibre_g_100"] = to_float(cell(row, "fibre_g"))  # for energy validation
            yield cursor, RawRecord(
                external_id=external_id,
                name=name,
                brand=None,                       # AFCD is generic foods, unbranded
                category=str(cell(row, "category") or "").strip() or None,
                source_url=self.SOURCE_URL,
                raw=raw,
                basis="per_100g",
                serving_size_desc="100 g",
                serving_weight_g=100.0,
                nutrients={
                    "energy_kj": to_float(cell(row, "energy_kj")),
                    "protein_g": to_float(cell(row, "protein_g")),
                    "carbs_g": to_float(cell(row, "carbs_g")),
                    "sugars_g": to_float(cell(row, "sugars_g")),
                    "fat_g": to_float(cell(row, "fat_g")),
                    "sat_fat_g": to_float(cell(row, "sat_fat_g")),
                    "sodium_mg": to_float(cell(row, "sodium_mg")),
                },
                source_version=self.source_version,
            )
